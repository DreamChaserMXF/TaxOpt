"""
TaxOpt FastAPI 后端 — 薄封装 main.py 的计算逻辑，供网页前端调用。
"""

import io
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional
from urllib.parse import quote

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, field_validator, model_validator

from main import (
    PRESETS_DIR,
    TaxCalculator,
    TaxConfig,
    TaxOptimizer,
    adaptive_search_step,
    load_tax_config_from_json,
    result_from_salary_bonus_split,
    validate_tax_config,
)

app = FastAPI(title="TaxOpt API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ──────────────────────────────────────────────
# 静态文件 / 前端入口
# ──────────────────────────────────────────────

WEB_DIR = Path(__file__).parent / "web"

if WEB_DIR.is_dir():
    app.mount("/static", StaticFiles(directory=WEB_DIR), name="static")


@app.get("/", include_in_schema=False)
def index():
    html = WEB_DIR / "index.html"
    if html.is_file():
        return FileResponse(html)
    return {"message": "TaxOpt API — web/index.html not found"}


# ──────────────────────────────────────────────
# Pydantic 模型
# ──────────────────────────────────────────────

class TaxConfigInput(BaseModel):
    basic_deduction: int = 60000
    social_security_base_min: float = 0
    social_security_base_limit: float = 0
    pension_rate: float = 0.08
    medical_rate: float = 0.02
    unemployment_rate: float = 0.005
    housing_fund_base_min: float = 0
    housing_fund_base_limit: float = 0
    housing_fund_employee_rate: float = 0.12
    housing_fund_employer_rate: float = 0.12
    children_education: float = 0
    continuing_education: float = 0
    serious_illness: float = 0
    housing_loan_interest: float = 0
    housing_rent: float = 0
    elderly_support: float = 0
    annual_additional_deduction: float = 0

    @field_validator("pension_rate", "medical_rate", "unemployment_rate",
                     "housing_fund_employee_rate", "housing_fund_employer_rate")
    @classmethod
    def rate_in_range(cls, v: float) -> float:
        if not 0 <= v <= 1:
            raise ValueError(f"费率须在 [0, 1] 区间内，当前为 {v}")
        return v

    @field_validator("children_education", "continuing_education", "serious_illness",
                     "housing_loan_interest", "housing_rent", "elderly_support",
                     "annual_additional_deduction", "basic_deduction")
    @classmethod
    def non_negative(cls, v: float) -> float:
        if v < 0:
            raise ValueError("扣除额不能为负数")
        return v


class CalculateRequest(BaseModel):
    mode: Literal["optimize", "salary", "bonus", "both"]
    total: Optional[float] = None
    monthly: Optional[float] = None
    bonus: Optional[float] = None
    objective: Literal["cash", "cash_plus_provident_fund"] = "cash"
    extra_income: float = 0.0
    stock_grants: List[float] = []
    config: Optional[TaxConfigInput] = None
    preset_slug: Optional[str] = None

    @model_validator(mode="after")
    def check_inputs(self):
        if self.config is None and self.preset_slug is None:
            raise ValueError("必须提供 config 或 preset_slug 之一")
        if self.mode == "salary" and self.monthly is None:
            raise ValueError("salary 模式需要 monthly")
        if self.mode == "salary" and self.total is None:
            raise ValueError("salary 模式需要 total")
        if self.mode == "bonus" and self.bonus is None:
            raise ValueError("bonus 模式需要 bonus")
        if self.mode == "bonus" and self.total is None:
            raise ValueError("bonus 模式需要 total")
        if self.mode == "both" and (self.monthly is None or self.bonus is None):
            raise ValueError("both 模式需要 monthly 和 bonus")
        if self.mode == "optimize" and self.total is None:
            raise ValueError("optimize 模式需要 total")
        return self


def _build_tax_config(req: CalculateRequest) -> TaxConfig:
    """preset_slug 优先；两者都有时 preset_slug 生效。"""
    if req.preset_slug:
        path = (PRESETS_DIR / f"{req.preset_slug}.json").resolve()
        if not path.is_file():
            raise HTTPException(status_code=404, detail=f"未知预设：{req.preset_slug}")
        return load_tax_config_from_json(path)
    ci = req.config
    return TaxConfig(
        basic_deduction=ci.basic_deduction,
        social_security_base_min=ci.social_security_base_min,
        social_security_base_limit=ci.social_security_base_limit,
        pension_rate=ci.pension_rate,
        medical_rate=ci.medical_rate,
        unemployment_rate=ci.unemployment_rate,
        housing_fund_base_min=ci.housing_fund_base_min,
        housing_fund_base_limit=ci.housing_fund_base_limit,
        housing_fund_employee_rate=ci.housing_fund_employee_rate,
        housing_fund_employer_rate=ci.housing_fund_employer_rate,
        children_education=ci.children_education,
        continuing_education=ci.continuing_education,
        serious_illness=ci.serious_illness,
        housing_loan_interest=ci.housing_loan_interest,
        housing_rent=ci.housing_rent,
        elderly_support=ci.elderly_support,
        annual_additional_deduction=ci.annual_additional_deduction,
    )


# ──────────────────────────────────────────────
# GET /api/presets
# ──────────────────────────────────────────────

_PROVINCE_LABEL = {
    "anhui": "安徽", "beijing": "北京", "chongqing": "重庆",
    "fujian": "福建", "gansu": "甘肃", "guangdong": "广东",
    "guangxi": "广西", "guizhou": "贵州", "hainan": "海南",
    "hebei": "河北", "heilongjiang": "黑龙江", "henan": "河南",
    "hubei": "湖北", "hunan": "湖南", "jiangsu": "江苏",
    "jiangxi": "江西", "jilin": "吉林", "liaoning": "辽宁",
    "neimenggu": "内蒙古", "ningxia": "宁夏", "qinghai": "青海",
    "shaanxi": "陕西", "shandong": "山东", "shanghai": "上海",
    "shanxi": "山西", "sichuan": "四川", "tianjin": "天津",
    "xinjiang": "新疆", "xizang": "西藏", "yunnan": "云南",
    "zhejiang": "浙江",
}

_CITY_LABEL = {
    "beijing": "北京", "shanghai": "上海", "tianjin": "天津",
    "chongqing": "重庆", "hangzhou": "杭州", "hangzhou-counties": "杭州（县域）",
    "ningbo": "宁波", "wenzhou": "温州", "nanjing": "南京",
    "suzhou": "苏州", "wuxi": "无锡", "guangzhou": "广州",
    "shenzhen": "深圳", "dongguan": "东莞", "foshan": "佛山",
    "fuzhou": "福州", "xiamen": "厦门", "jinan": "济南",
    "qingdao": "青岛", "wuhan": "武汉", "changsha": "长沙",
    "zhengzhou": "郑州", "hefei": "合肥", "nanchang": "南昌",
    "taiyuan": "太原", "xian": "西安", "shenyang": "沈阳",
    "dalian": "大连", "changchun": "长春", "haerbin": "哈尔滨",
    "kunming": "昆明", "guiyang": "贵阳", "nanning": "南宁",
    "lanzhou": "兰州", "haikou": "海口", "yinchuan": "银川",
    "xining": "西宁", "wulumuqi": "乌鲁木齐", "huhehaote": "呼和浩特",
    "lasa": "拉萨", "chengdu": "成都", "shijiazhuang": "石家庄",
}


@app.get("/api/presets")
def list_presets():
    presets = []
    for p in sorted(PRESETS_DIR.glob("*.json")):
        slug = p.stem
        parts = slug.split("-", 1)
        province = _PROVINCE_LABEL.get(parts[0], parts[0])
        city = _CITY_LABEL.get(parts[1], parts[1]) if len(parts) > 1 else ""
        presets.append({"slug": slug, "label": f"{province} · {city}", "province": province})
    return presets


@app.get("/api/presets/{slug}")
def get_preset(slug: str):
    path = (PRESETS_DIR / f"{slug}.json").resolve()
    if not path.is_file():
        raise HTTPException(status_code=404, detail=f"未知预设：{slug}")
    cfg = load_tax_config_from_json(path)
    return {
        "basic_deduction": cfg.basic_deduction,
        "social_security_base_min": cfg.social_security_base_min,
        "social_security_base_limit": cfg.social_security_base_limit,
        "pension_rate": cfg.pension_rate,
        "medical_rate": cfg.medical_rate,
        "unemployment_rate": cfg.unemployment_rate,
        "housing_fund_base_min": cfg.housing_fund_base_min,
        "housing_fund_base_limit": cfg.housing_fund_base_limit,
        "housing_fund_employee_rate": cfg.housing_fund_employee_rate,
        "housing_fund_employer_rate": cfg.housing_fund_employer_rate,
        "children_education": cfg.children_education,
        "continuing_education": cfg.continuing_education,
        "serious_illness": cfg.serious_illness,
        "housing_loan_interest": cfg.housing_loan_interest,
        "housing_rent": cfg.housing_rent,
        "elderly_support": cfg.elderly_support,
        "annual_additional_deduction": cfg.annual_additional_deduction,
    }


# ──────────────────────────────────────────────
# POST /api/calculate
# ──────────────────────────────────────────────

@app.post("/api/calculate")
def calculate(req: CalculateRequest):
    try:
        cfg = _build_tax_config(req)
        validate_tax_config(cfg)
    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    calc = TaxCalculator(cfg)

    try:
        if req.mode == "optimize":
            total = req.total
            step = adaptive_search_step(total)
            result = TaxOptimizer(calc).optimize(total, step, objective=req.objective)
            if not result:
                raise HTTPException(status_code=422, detail="未找到优化方案")

        elif req.mode == "salary":
            total = req.total
            monthly = req.monthly
            bonus = total - monthly * 12
            if bonus < -0.01:
                raise HTTPException(status_code=422, detail="月薪×12 超过全年收入")
            result = result_from_salary_bonus_split(
                calc, total, monthly, max(bonus, 0.0),
                extra_income=req.extra_income, stock_grants=req.stock_grants,
            )

        elif req.mode == "bonus":
            total = req.total
            bonus = req.bonus
            monthly = (total - bonus) / 12.0
            if monthly < -0.01:
                raise HTTPException(status_code=422, detail="年终奖超过全年收入")
            result = result_from_salary_bonus_split(
                calc, total, max(monthly, 0.0), bonus,
                extra_income=req.extra_income, stock_grants=req.stock_grants,
            )

        else:  # both
            monthly = req.monthly
            bonus = req.bonus
            total = monthly * 12 + bonus
            result = result_from_salary_bonus_split(
                calc, total, monthly, bonus,
                extra_income=req.extra_income, stock_grants=req.stock_grants,
            )

    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    return result


# ──────────────────────────────────────────────
# POST /api/export/excel
# ──────────────────────────────────────────────

class ExportRequest(BaseModel):
    result: Dict[str, Any]
    title: Optional[str] = None


def _build_excel(result: Dict[str, Any], title: str) -> io.BytesIO:
    wb = Workbook()

    # ── 汇总页 ────────────────────────────────
    ws = wb.active
    ws.title = "汇总"

    header_font  = Font(bold=True, size=12, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="4F46E5")
    center       = Alignment(horizontal="center", vertical="center")
    label_font   = Font(bold=True)

    # 标题行
    ws.merge_cells("A1:C1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 28

    ws.append([])

    # 列头
    ws.append(["指标", "金额 / 比率"])
    for cell in ws[ws.max_row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    nominal = (result.get("annual_salary") or 0) + (result.get("bonus") or 0)
    rows = [
        ("名义收入",           nominal),
        ("到手现金",           result.get("net_take_home")),
        ("公积金入账（个人+单位）", result.get("annual_provident_fund")),
        ("现金 + 公积金",      result.get("net_take_home_including_provident_fund")),
        ("全年个税",           result.get("total_tax")),
        ("税率",               f'{result.get("effective_tax_rate", 0)*100:.2f}%'),
        ("个人社保（三险）",    result.get("annual_insurance")),
        ("广义税率（个税+社保）", f'{result.get("effective_burden_rate", 0)*100:.2f}%'),
    ]
    if result.get("bonus"):
        rows += [
            ("年终奖",          result.get("bonus")),
            ("年终奖个税",       result.get("bonus_tax")),
            ("年终奖税后到手",   result.get("bonus_after_tax")),
        ]

    for label, value in rows:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = label_font

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20

    # ── 月度明细页 ────────────────────────────
    ws2 = wb.create_sheet("月度明细")
    detail_headers = ["月份", "税前月薪", "个人三险", "个人公积金", "单位公积金", "当月税额", "税后到手", "广义税后到手"]
    ws2.append(detail_headers)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row in result.get("monthly_details", []):
        ins = row.get("insurance", {})
        hf  = row.get("housing_fund", {})
        ws2.append([
            f'{row["month"]}月',
            row.get("salary"),
            ins.get("total"),
            hf.get("employee"),
            hf.get("employer"),
            row.get("tax"),
            row.get("after_tax_salary"),
            row.get("after_tax_including_provident_fund"),
        ])

    for col in ["A","B","C","D","E","F","G","H"]:
        ws2.column_dimensions[col].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.post("/api/export/excel")
def export_excel(req: ExportRequest):
    nominal = (req.result.get("annual_salary") or 0) + (req.result.get("bonus") or 0)
    title = req.title or f"个税税筹结果（名义收入 {nominal:,.0f} 元）"
    buf = _build_excel(req.result, title)
    filename = title + ".xlsx"
    encoded = quote(filename, safe="")
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"}
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
